"""Extract tutorial process data from course materials (adapted from tutorial-to-notes Phase 1)."""

from __future__ import annotations

import re
import uuid
from pathlib import Path
from typing import Any

MIN_IMAGE_BYTES = 5 * 1024


def extract_material_process(
    *,
    course_id: str,
    material_id: str | None = None,
    chapter_key: str | None = None,
    work_root: Path,
) -> dict[str, Any]:
    from app.core.dependencies import get_course_catalog

    catalog = get_course_catalog()
    if catalog.get_course(course_id) is None:
        raise ValueError(f"课程不存在: {course_id}")

    targets: list[Any] = []
    if material_id:
        material = catalog.get_material(material_id)
        if material is None or material.course_id != course_id:
            raise ValueError(f"未找到 material_id={material_id}")
        targets = [material]
    elif chapter_key:
        key = str(chapter_key).strip()
        targets = [m for m in catalog.list_materials(course_id) if m.chapter_key == key]
        if not targets:
            raise ValueError(f"章节 {chapter_key} 下无课件文件")
    else:
        raise ValueError("请指定 material_id 或 chapter_key")

    job_id = uuid.uuid4().hex[:10]
    job_dir = work_root / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    all_pages: list[dict[str, Any]] = []
    documents: list[dict[str, Any]] = []
    page_offset = 0

    for material in sorted(targets, key=lambda m: m.title):
        suffix = material.path.suffix.lower()
        if suffix == ".pdf":
            pages = _extract_pdf(material.path, job_dir / material.id)
        elif suffix == ".pptx":
            pages = _extract_pptx(material.path, job_dir / material.id)
        else:
            pages = _extract_plain_text(material.path)

        for page in pages:
            page["page"] = page_offset + int(page["page"])
            page["material_id"] = material.id
            page["material_title"] = material.title
        page_offset += len(pages)
        all_pages.extend(pages)
        documents.append(
            {
                "material_id": material.id,
                "title": material.title,
                "chapter_key": material.chapter_key,
                "page_count": len(pages),
                "doc_type": material.doc_type,
            }
        )

    if not all_pages:
        raise ValueError("未能从资料中提取有效页面")

    _mark_redundant_images(all_pages)
    pending_images = sum(
        1 for p in all_pages for img in p.get("images", []) if img.get("status") == "待理解"
    )

    return {
        "job_id": job_id,
        "course_id": course_id,
        "chapter_key": chapter_key,
        "material_id": material_id,
        "documents": documents,
        "pages": all_pages,
        "page_count": len(all_pages),
        "pending_images": pending_images,
        "summary": f"已提取 {len(documents)} 个文件、{len(all_pages)} 页，待视觉理解 {pending_images} 张图",
    }


def _extract_pdf(path: Path, image_dir: Path) -> list[dict[str, Any]]:
    import fitz

    image_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(path)
    pages: list[dict[str, Any]] = []
    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        images: list[dict[str, Any]] = []
        for img_idx, img_info in enumerate(page.get_images(full=True), start=1):
            xref = img_info[0]
            try:
                extracted = doc.extract_image(xref)
                blob = extracted.get("image") or b""
                if len(blob) < MIN_IMAGE_BYTES:
                    continue
                ext = str(extracted.get("ext") or "png")
                img_path = image_dir / f"p{idx}_img{img_idx}.{ext}"
                img_path.write_bytes(blob)
                images.append(
                    {
                        "id": f"p{idx}_img{img_idx}",
                        "path": str(img_path),
                        "size_bytes": len(blob),
                        "status": "待理解",
                    }
                )
            except Exception:
                continue
        pages.append({"page": idx, "text": text, "images": images})
    return pages


def _extract_pptx(path: Path, image_dir: Path) -> list[dict[str, Any]]:
    from pptx import Presentation

    image_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(path)
    pages: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        images: list[dict[str, Any]] = []
        img_idx = 0
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                line = shape.text.strip()
                if line:
                    lines.append(line)
            if not hasattr(shape, "image"):
                continue
            try:
                blob = shape.image.blob
                if len(blob) < MIN_IMAGE_BYTES:
                    continue
                img_idx += 1
                ext = str(getattr(shape.image, "ext", None) or "png")
                img_path = image_dir / f"p{idx}_img{img_idx}.{ext}"
                img_path.write_bytes(blob)
                images.append(
                    {
                        "id": f"p{idx}_img{img_idx}",
                        "path": str(img_path),
                        "size_bytes": len(blob),
                        "status": "待理解",
                    }
                )
            except Exception:
                continue
        pages.append({"page": idx, "text": "\n".join(lines), "images": images})
    return pages


def _extract_plain_text(path: Path) -> list[dict[str, Any]]:
    from app.infrastructure.document_parser import extract_text

    text, _ = extract_text(path)
    chunks = re.split(r"^## 第\s*(\d+)\s*页\s*$", text, flags=re.MULTILINE)
    pages: list[dict[str, Any]] = []
    if len(chunks) > 2:
        i = 1
        while i < len(chunks) - 1:
            page_num = int(chunks[i])
            body = chunks[i + 1].strip()
            pages.append({"page": page_num, "text": body, "images": []})
            i += 2
    else:
        pages.append({"page": 1, "text": text.strip(), "images": []})
    return pages


def _text_overlap_ratio(image_text: str, page_text: str) -> float:
    if not image_text.strip() or not page_text.strip():
        return 0.0
    img_tokens = {t for t in re.split(r"[\s\|，。；、]+", image_text) if len(t) >= 2}
    if not img_tokens:
        return 0.0
    page_blob = page_text
    hit = sum(1 for t in img_tokens if t in page_blob)
    return hit / len(img_tokens)


def _mark_redundant_images(pages: list[dict[str, Any]]) -> None:
    """If page text is rich, skip decorative images; vision overlap handled later."""
    for page in pages:
        text = str(page.get("text") or "")
        for img in page.get("images") or []:
            if len(text) > 50 and img.get("size_bytes", 0) < 10 * 1024:
                img["status"] = "跳过（装饰性）"


def apply_vision_to_process(pages: list[dict[str, Any]], image_id: str, analysis: str) -> dict[str, Any] | None:
    parsed = _parse_vision_output(analysis)
    for page in pages:
        for img in page.get("images") or []:
            if img.get("id") != image_id:
                continue
            page_text = str(page.get("text") or "")
            overlap = _text_overlap_ratio(parsed.get("description", ""), page_text)
            if overlap >= 0.8:
                img["status"] = "跳过（与页内文本重复）"
                img["keep"] = "C"
            else:
                img["status"] = "已理解"
                img["vision_type"] = parsed.get("type", "")
                img["description"] = parsed.get("description", "")
                img["keep"] = parsed.get("keep", "B")
            return img
    return None


def _parse_vision_output(text: str) -> dict[str, str]:
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    result = {"type": "", "description": "", "keep": "B"}
    desc_parts: list[str] = []
    for line in lines:
        if line.startswith("类型:") or line.startswith("类型："):
            result["type"] = line.split(":", 1)[-1].split("：", 1)[-1].strip()
        elif line.startswith("内容描述:") or line.startswith("内容描述："):
            desc_parts.append(line.split(":", 1)[-1].split("：", 1)[-1].strip())
        elif line.startswith("保留判定:") or line.startswith("保留判定："):
            keep_raw = line.split(":", 1)[-1].split("：", 1)[-1].strip().upper()
            result["keep"] = keep_raw[:1] if keep_raw else "B"
        elif not line.startswith("类型") and not line.startswith("保留判定"):
            desc_parts.append(line)
    if desc_parts:
        result["description"] = " ".join(desc_parts)
    return result


def serialize_process_for_llm(process: dict[str, Any], *, max_pages: int | None = None) -> str:
    lines: list[str] = []
    lines.append(process.get("summary", ""))
    structure = process.get("structure_report")
    if structure:
        lines.append("\n## 章节结构\n" + structure)

    pages = process.get("pages") or []
    if max_pages is not None:
        pages = pages[:max_pages]

    for page in pages:
        pnum = page.get("page")
        title = page.get("material_title", "")
        text = str(page.get("text") or "").strip()
        lines.append(f"\n## 第{pnum}页 {title}")
        if text:
            lines.append("### 文本\n" + text[:2000])
        for img in page.get("images") or []:
            if img.get("status", "").startswith("跳过"):
                continue
            src_name = Path(str(img.get("path") or "")).name
            desc = img.get("description") or img.get("status", "待理解")
            keep = img.get("keep", "")
            asset_hint = f"images/{src_name}" if src_name else ""
            hint = f" [引用={asset_hint}]" if asset_hint else ""
            lines.append(f"### 图片 {img.get('id')} [keep={keep}]{hint}\n{desc}")
    return "\n".join(lines)
