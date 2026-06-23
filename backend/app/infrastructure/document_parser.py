from pathlib import Path


def extract_text(path: Path) -> tuple[str, int | None]:
    """Extract readable text from supported course material files."""
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return text, None
    if suffix == ".pdf":
        return _extract_pdf(path)
    return f"暂不支持在线预览 {path.name}，请下载原文件阅读。", None


def _extract_pptx(path: Path) -> tuple[str, int | None]:
    from pptx import Presentation

    prs = Presentation(path)
    sections: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                line = shape.text.strip()
                if line:
                    lines.append(line)
        if lines:
            sections.append(f"## 第 {idx} 页\n\n" + "\n\n".join(lines))
    if not sections:
        return f"PPT 文件 `{path.name}` 未提取到文本内容。", len(prs.slides)
    return f"# {path.stem}\n\n" + "\n\n".join(sections), len(prs.slides)


def _extract_docx(path: Path) -> tuple[str, int | None]:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        return f"Word 文件 `{path.name}` 未提取到文本内容。", None
    return f"# {path.stem}\n\n" + "\n\n".join(paragraphs), None


def _extract_pdf(path: Path) -> tuple[str, int | None]:
    try:
        import fitz
    except ImportError:
        return f"PDF 文件 `{path.name}` 需安装 PyMuPDF 后预览，请下载原文件阅读。", None

    doc = fitz.open(path)
    pages: list[str] = []
    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if text:
            pages.append(f"## 第 {idx} 页\n\n{text}")
    if not pages:
        return f"PDF 文件 `{path.name}` 未提取到文本内容。", len(doc)
    return f"# {path.stem}\n\n" + "\n\n".join(pages), len(doc)
